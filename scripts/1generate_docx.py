import os
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Ajusta márgenes internos de las celdas de tabla."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def add_heading_with_bottom_border(doc, text):
    """Agrega un encabezado de sección elegante con línea inferior en tono azul corporativo."""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(text.upper())
    run.font.name = 'Calibri'
    run.font.size = Pt(11)
    run.font.bold = True
    run.font.color.rgb = RGBColor(31, 78, 121)  # Azul corporativo #1F4E79

    # Añadir borde inferior
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), '6')  # Grosor de la línea
    bottom.set(qn('w:space'), '1')
    bottom.set(qn('w:color'), '1F4E79')
    pBdr.append(bottom)
    pPr.append(pBdr)

def render_exact_cv_layout(datos: dict, output_filename: str):
    """
    Genera el archivo Word (.docx) procesando la estructura de datos extraída del CV.
    """
    doc = Document()

    # Configuración de márgenes estándar de 2.0 cm
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # --- ENCABEZADO (DATOS PERSONALES) ---
    p_nombre = doc.add_paragraph()
    p_nombre.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_nombre.paragraph_format.space_after = Pt(2)
    run_nombre = p_nombre.add_run(datos.get("nombre", "NOMBRE COMPLETO").upper())
    run_nombre.font.name = 'Calibri'
    run_nombre.font.size = Pt(20)
    run_nombre.font.bold = True
    run_nombre.font.color.rgb = RGBColor(31, 78, 121)

    if datos.get("titulo"):
        p_titulo = doc.add_paragraph()
        p_titulo.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_titulo.paragraph_format.space_after = Pt(4)
        run_titulo = p_titulo.add_run(datos.get("titulo"))
        run_titulo.font.name = 'Calibri'
        run_titulo.font.size = Pt(12)
        run_titulo.font.italic = True
        run_titulo.font.color.rgb = RGBColor(89, 89, 89)

    # Datos de contacto en una sola línea
    contact_info = []
    if datos.get("telefono"): contact_info.append(f"📞 {datos['telefono']}")
    if datos.get("email"): contact_info.append(f"📧 {datos['email']}")
    if datos.get("ubicacion"): contact_info.append(f"📍 {datos['ubicacion']}")

    if contact_info:
        p_contacto = doc.add_paragraph()
        p_contacto.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_contacto.paragraph_format.space_after = Pt(12)
        run_contacto = p_contacto.add_run("  |  ".join(contact_info))
        run_contacto.font.name = 'Calibri'
        run_contacto.font.size = Pt(9.5)
        run_contacto.font.color.rgb = RGBColor(51, 51, 51)

    # --- PERFIL PROFESIONAL ---
    if datos.get("perfil"):
        add_heading_with_bottom_border(doc, "Perfil Profesional")
        p_perfil = doc.add_paragraph()
        p_perfil.paragraph_format.space_after = Pt(8)
        p_perfil.paragraph_format.line_spacing = 1.15
        run_perfil = p_perfil.add_run(datos["perfil"])
        run_perfil.font.name = 'Calibri'
        run_perfil.font.size = Pt(10.5)

    # --- EXPERIENCIA LABORAL ---
    experiencias = datos.get("experiencia", [])
    if experiencias:
        add_heading_with_bottom_border(doc, "Experiencia Laboral")
        for exp in experiencias:
            p_exp = doc.add_paragraph()
            p_exp.paragraph_format.space_before = Pt(4)
            p_exp.paragraph_format.space_after = Pt(2)
            
            # Puesto / Cargo
            run_cargo = p_exp.add_run(exp.get("cargo", "Puesto Desempeñado"))
            run_cargo.font.name = 'Calibri'
            run_cargo.font.size = Pt(11)
            run_cargo.font.bold = True
            
            # Empresa
            if exp.get("empresa"):
                run_empresa = p_exp.add_run(f" — {exp['empresa']}")
                run_empresa.font.name = 'Calibri'
                run_empresa.font.size = Pt(10.5)
                run_empresa.font.italic = True

            # Responsabilidades o logros
            responsabilidades = exp.get("responsabilidades", [])
            for resp in responsabilidades:
                p_resp = doc.add_paragraph(style='List Bullet')
                p_resp.paragraph_format.space_after = Pt(2)
                run_resp = p_resp.add_run(resp)
                run_resp.font.name = 'Calibri'
                run_resp.font.size = Pt(10)

    # --- EDUCACIÓN ---
    educacion = datos.get("educacion", [])
    if educacion:
        add_heading_with_bottom_border(doc, "Educación")
        for edu in educacion:
            p_edu = doc.add_paragraph()
            p_edu.paragraph_format.space_before = Pt(3)
            p_edu.paragraph_format.space_after = Pt(2)
            
            run_edu = p_edu.add_run(edu.get("titulo", "Grado Académico"))
            run_edu.font.name = 'Calibri'
            run_edu.font.size = Pt(10.5)
            run_edu.font.bold = True

            inst = []
            if edu.get("institucion"): inst.append(edu["institucion"])
            if edu.get("estado"): inst.append(edu["estado"])

            if inst:
                run_inst = p_edu.add_run(f" | {', '.join(inst)}")
                run_inst.font.name = 'Calibri'
                run_inst.font.size = Pt(10)

    # --- HABILIDADES ---
    habilidades = datos.get("habilidades", [])
    if habilidades:
        add_heading_with_bottom_border(doc, "Habilidades y Competencias")
        p_hab = doc.add_paragraph()
        p_hab.paragraph_format.space_before = Pt(4)
        p_hab.paragraph_format.space_after = Pt(6)
        
        hab_texto = "  •  ".join(habilidades)
        run_hab = p_hab.add_run(hab_texto)
        run_hab.font.name = 'Calibri'
        run_hab.font.size = Pt(10)

    # Guardar documento
    doc.save(output_filename)

if __name__ == "__main__":
    pass