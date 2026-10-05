import os
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def set_cell_background(cell, hex_color):
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=140, bottom=140, left=160, right=160):
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def remove_table_borders(table):
    tblPr = table._element.xpath('w:tblPr')
    if tblPr:
        borders = parse_xml(
            f'<w:tblBorders {nsdecls("w")}>'
            f'<w:top w:val="none"/>'
            f'<w:left w:val="none"/>'
            f'<w:bottom w:val="none"/>'
            f'<w:right w:val="none"/>'
            f'<w:insideH w:val="none"/>'
            f'<w:insideV w:val="none"/>'
            f'</w:tblBorders>'
        )
        tblPr[0].append(borders)

def add_heading_main(doc_or_cell, text):
    p = doc_or_cell.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(text.upper())
    run.font.name = 'Calibri'
    run.font.size = Pt(12)
    run.font.bold = True
    run.font.color.rgb = RGBColor(15, 32, 67)

    pPr = p._p.get_or_add_pPr()
    pBdr = parse_xml(
        f'<w:pBdr {nsdecls("w")}>'
        f'<w:bottom w:val="single" w:sz="12" w:space="3" w:color="0F2043"/>'
        f'</w:pBdr>'
    )
    pPr.append(pBdr)

def add_heading_sidebar(cell, text):
    p = cell.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(text.upper())
    run.font.name = 'Calibri'
    run.font.size = Pt(11)
    run.font.bold = True
    run.font.color.rgb = RGBColor(15, 32, 67)

    pPr = p._p.get_or_add_pPr()
    pBdr = parse_xml(
        f'<w:pBdr {nsdecls("w")}>'
        f'<w:bottom w:val="single" w:sz="8" w:space="3" w:color="0F2043"/>'
        f'</w:pBdr>'
    )
    pPr.append(pBdr)

def render_exact_cv_layout(datos: dict, output_filename: str):
    doc = Document()

    for section in doc.sections:
        section.top_margin = Inches(0.35)
        section.bottom_margin = Inches(0.35)
        section.left_margin = Inches(0.35)
        section.right_margin = Inches(0.35)

    # BANNER SUPERIOR
    tbl_hdr = doc.add_table(rows=1, cols=1)
    tbl_hdr.alignment = WD_TABLE_ALIGNMENT.CENTER
    remove_table_borders(tbl_hdr)
    c_hdr = tbl_hdr.cell(0, 0)
    c_hdr.width = Inches(7.6)
    set_cell_background(c_hdr, "0F2043")
    set_cell_margins(c_hdr, top=180, bottom=180, left=200, right=200)

    p_nom = c_hdr.paragraphs[0]
    p_nom.paragraph_format.space_after = Pt(2)
    r_nom = p_nom.add_run(datos.get("nombre", "NOMBRE COMPLETO").upper())
    r_nom.font.name = 'Calibri'
    r_nom.font.size = Pt(24)
    r_nom.font.bold = True
    r_nom.font.color.rgb = RGBColor(255, 255, 255)

    if datos.get("titulo"):
        p_tit = c_hdr.add_paragraph()
        p_tit.paragraph_format.space_after = Pt(0)
        r_tit = p_tit.add_run(datos.get("titulo"))
        r_tit.font.name = 'Calibri'
        r_tit.font.size = Pt(13)
        r_tit.font.italic = True
        r_tit.font.color.rgb = RGBColor(190, 215, 240)

    p_sp = doc.add_paragraph()
    p_sp.paragraph_format.space_before = Pt(4)
    p_sp.paragraph_format.space_after = Pt(0)

    # MAQUETACIÓN CUERPO
    tbl_body = doc.add_table(rows=1, cols=2)
    tbl_body.alignment = WD_TABLE_ALIGNMENT.CENTER
    remove_table_borders(tbl_body)

    c_left = tbl_body.cell(0, 0)
    c_right = tbl_body.cell(0, 1)

    c_left.width = Inches(2.6)
    c_right.width = Inches(5.0)

    set_cell_background(c_left, "F4F6F9")
    set_cell_margins(c_left, top=140, bottom=140, left=140, right=140)
    set_cell_margins(c_right, top=140, bottom=140, left=180, right=100)

    # COLUMNA IZQUIERDA
    c_left.paragraphs[0].text = ""

    foto_path = datos.get("foto_path")
    if foto_path and os.path.exists(foto_path):
        p_foto = c_left.add_paragraph()
        p_foto.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_foto.paragraph_format.space_before = Pt(2)
        p_foto.paragraph_format.space_after = Pt(6)
        run_foto = p_foto.add_run()
        run_foto.add_picture(foto_path, width=Inches(1.8))

    add_heading_sidebar(c_left, "Contacto")
    contacto_items = []
    if datos.get("telefono"): contacto_items.append(f"📱  {datos['telefono']}")
    if datos.get("email"): contacto_items.append(f"✉  {datos['email']}")
    if datos.get("linkedin"): contacto_items.append(f"🔗  {datos['linkedin']}")
    if datos.get("ubicacion"): contacto_items.append(f"📍  {datos['ubicacion']}")

    for item in contacto_items:
        p_c = c_left.add_paragraph()
        p_c.paragraph_format.space_after = Pt(3)
        p_c.paragraph_format.line_spacing = 1.1
        r_c = p_c.add_run(item)
        r_c.font.name = 'Calibri'
        r_c.font.size = Pt(9.5)
        r_c.font.color.rgb = RGBColor(40, 40, 40)

    habilidades = datos.get("habilidades", [])
    if habilidades:
        add_heading_sidebar(c_left, "Habilidades")
        for hab in habilidades:
            p_h = c_left.add_paragraph()
            p_h.paragraph_format.space_after = Pt(2)
            r_h = p_h.add_run(f"▪  {hab}")
            r_h.font.name = 'Calibri'
            r_h.font.size = Pt(10)
            r_h.font.color.rgb = RGBColor(40, 40, 40)

    idiomas = datos.get("idiomas", [])
    if idiomas:
        add_heading_sidebar(c_left, "Idiomas")
        for idioma in idiomas:
            p_lang = c_left.add_paragraph()
            p_lang.paragraph_format.space_after = Pt(2)
            r_l = p_lang.add_run(idioma.get("nombre", ""))
            r_l.font.name = 'Calibri'
            r_l.font.size = Pt(10)
            r_l.font.bold = True
            r_l.font.color.rgb = RGBColor(15, 32, 67)

            if idioma.get("nivel"):
                r_lvl = p_lang.add_run(f" — {idioma['nivel']}")
                r_lvl.font.name = 'Calibri'
                r_lvl.font.size = Pt(9.5)
                r_lvl.font.color.rgb = RGBColor(60, 60, 60)

    # COLUMNA DERECHA
    c_right.paragraphs[0].text = ""

    if datos.get("perfil"):
        add_heading_main(c_right, "Perfil Profesional")
        p_p = c_right.add_paragraph()
        p_p.paragraph_format.space_after = Pt(6)
        p_p.paragraph_format.line_spacing = 1.15
        r_p = p_p.add_run(datos["perfil"])
        r_p.font.name = 'Calibri'
        r_p.font.size = Pt(10.5)
        r_p.font.color.rgb = RGBColor(35, 35, 35)

    educacion = datos.get("educacion", [])
    if educacion:
        add_heading_main(c_right, "Educación")
        for edu in educacion:
            p_e = c_right.add_paragraph()
            p_e.paragraph_format.space_before = Pt(4)
            p_e.paragraph_format.space_after = Pt(2)

            r_e = p_e.add_run(edu.get("titulo", ""))
            r_e.font.name = 'Calibri'
            r_e.font.size = Pt(11)
            r_e.font.bold = True
            r_e.font.color.rgb = RGBColor(15, 32, 67)

            inst = []
            if edu.get("institucion"): inst.append(edu["institucion"])
            if edu.get("estado"): inst.append(str(edu["estado"]))
            if inst:
                r_i = p_e.add_run(f" | {' - '.join(inst)}")
                r_i.font.name = 'Calibri'
                r_i.font.size = Pt(10)
                r_i.font.italic = True
                r_i.font.color.rgb = RGBColor(90, 90, 90)

            if edu.get("detalle"):
                p_d = c_right.add_paragraph()
                p_d.paragraph_format.left_indent = Inches(0.12)
                p_d.paragraph_format.space_after = Pt(4)
                r_d = p_d.add_run(f"•  {edu['detalle']}")
                r_d.font.name = 'Calibri'
                r_d.font.size = Pt(9.5)
                r_d.font.color.rgb = RGBColor(60, 60, 60)

    experiencias = datos.get("experiencia", [])
    if experiencias:
        add_heading_main(c_right, "Experiencia Laboral")
        for exp in experiencias:
            p_exp = c_right.add_paragraph()
            p_exp.paragraph_format.space_before = Pt(4)
            p_exp.paragraph_format.space_after = Pt(2)

            r_cargo = p_exp.add_run(exp.get("cargo", "Puesto"))
            r_cargo.font.name = 'Calibri'
            r_cargo.font.size = Pt(11.5)
            r_cargo.font.bold = True
            r_cargo.font.color.rgb = RGBColor(15, 32, 67)

            if exp.get("empresa"):
                r_emp = p_exp.add_run(f" | {exp['empresa']}")
                r_emp.font.name = 'Calibri'
                r_emp.font.size = Pt(10.5)
                r_emp.font.italic = True
                r_emp.font.color.rgb = RGBColor(90, 90, 90)

            if exp.get("periodo"):
                p_per = c_right.add_paragraph()
                p_per.paragraph_format.space_after = Pt(3)
                r_per = p_per.add_run(f"🗓 {exp['periodo']}")
                r_per.font.name = 'Calibri'
                r_per.font.size = Pt(9.5)
                r_per.font.bold = True
                r_per.font.color.rgb = RGBColor(120, 120, 120)

            for resp in exp.get("responsabilidades", []):
                p_r = c_right.add_paragraph()
                p_r.paragraph_format.left_indent = Inches(0.12)
                p_r.paragraph_format.space_after = Pt(2)
                p_r.paragraph_format.line_spacing = 1.15
                r_r = p_r.add_run(f"•  {resp}")
                r_r.font.name = 'Calibri'
                r_r.font.size = Pt(10)
                r_r.font.color.rgb = RGBColor(45, 45, 45)

    if datos.get("informacion_adicional"):
        add_heading_main(c_right, "Información Adicional")
        p_info = c_right.add_paragraph()
        p_info.paragraph_format.space_after = Pt(4)
        r_info = p_info.add_run(datos["informacion_adicional"])
        r_info.font.name = 'Calibri'
        r_info.font.size = Pt(10)
        r_info.font.color.rgb = RGBColor(35, 35, 35)

    if datos.get("referencias"):
        add_heading_main(c_right, "Referencias")
        p_ref = c_right.add_paragraph()
        p_ref.paragraph_format.space_after = Pt(4)
        r_ref = p_ref.add_run(datos["referencias"])
        r_ref.font.name = 'Calibri'
        r_ref.font.size = Pt(10)
        r_ref.font.color.rgb = RGBColor(35, 35, 35)

    os.makedirs(os.path.dirname(output_filename), exist_ok=True)
    doc.save(output_filename)